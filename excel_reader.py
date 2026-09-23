"""
Excel Reader Module

Reads route-order workbooks and exposes a clean route table for the renamer.
The reader supports both ordinary single-header sheets and the attached style
where one sheet contains an A-to-B table followed by a B-to-A table.
"""

from __future__ import annotations

import os
import re
from typing import Dict, List, Optional, Tuple

import pandas as pd
from openpyxl import load_workbook


class ExcelReader:

    COLUMN_VARIATIONS = {
        "serial": [
            "Serial",
            "Sr No",
            "Sr. No",
            "Sequence",
            "Sequence No",
            "Serial No",
            "Serial Number",
        ],
        "start_point": [
            "Start Point",
            "Start",
            "Start Location",
            "From",
            "Origin",
            "Landmark From",
            "Landmark (From)",
        ],
        "end_point": [
            "End Point",
            "End",
            "End Location",
            "To",
            "Destination",
            "Landmark To",
            "Landmark (To)",
        ],
        "road_type": [
            "Road Type",
            "Type",
            "Road",
            "Surface Type",
            "Type of Road",
        ],
        "distance": [
            "Total Distance of Each Section",
            "Total Distance",
            "Section Distance",
            "Distance / km",
            "Distance (km)",
            "Distance Covered (km)",
            "Distance Covered",
            "Distance",
            "KM",
            "Length",
        ],
    }

    REQUIRED_FIELDS = ["start_point", "end_point", "road_type"]

    def __init__(self):
        self.df: Optional[pd.DataFrame] = None
        self.raw_df: Optional[pd.DataFrame] = None
        self.file_path: Optional[str] = None
        self.column_mapping: Dict[str, str] = {}

    def read_file(self, file_path: str) -> Tuple[bool, str]:
        """Read an Excel file and auto-detect route columns."""
        try:
            if not os.path.exists(file_path):
                return False, f"File not found: {file_path}"

            self.file_path = file_path
            parsed_rows = self._read_route_sections(file_path)

            if parsed_rows:
                self.df = pd.DataFrame(parsed_rows)
                self.raw_df = self.df.copy()
                self.column_mapping = {
                    "serial": "serial",
                    "start_point": "start_point",
                    "end_point": "end_point",
                    "road_type": "road_type",
                }
                if "distance" in self.df.columns:
                    self.column_mapping["distance"] = "distance"
                return True, f"Successfully read {len(self.df)} route rows from Excel"

            self.raw_df = pd.read_excel(file_path)
            self.df = self.raw_df.copy()
            self.column_mapping = self._auto_detect_columns()
            if not self.column_mapping:
                return False, "Could not auto-detect route columns in Excel file"

            return True, f"Successfully read {len(self.df)} rows from Excel"
        except Exception as exc:
            return False, f"Error reading Excel file: {exc}"

    def _read_route_sections(self, file_path: str) -> List[Dict]:
        """Parse workbooks with one or more route tables per sheet."""
        workbook = load_workbook(file_path, data_only=True)
        routes: List[Dict] = []
        serial_counter = 1

        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            row_index = 0
            section_number = 0

            while row_index < len(rows):
                header = rows[row_index]
                mapping = self._detect_header_mapping(header)

                if not self._has_route_header(mapping):
                    row_index += 1
                    continue

                section_number += 1
                row_index += 1

                while row_index < len(rows):
                    values = rows[row_index]
                    next_mapping = self._detect_header_mapping(values)
                    if self._has_route_header(next_mapping):
                        break

                    if self._is_blank_row(values):
                        row_index += 1
                        continue

                    row = self._row_from_mapping(
                        values=values,
                        mapping=mapping,
                        serial_counter=serial_counter,
                        sheet_name=sheet.title,
                        source_excel_row=row_index + 1,
                        section_number=section_number,
                    )
                    if row:
                        routes.append(row)
                        serial_counter += 1

                    row_index += 1

        return routes

    def _detect_header_mapping(self, row_values) -> Dict[str, int]:
        mapping: Dict[str, int] = {}
        normalized_cells = [self._normalize_header(value) for value in row_values]

        for idx, normalized in enumerate(normalized_cells):
            if not normalized:
                continue

            for field, variations in self.COLUMN_VARIATIONS.items():
                if field in mapping:
                    continue
                for variation in variations:
                    if normalized == self._normalize_header(variation):
                        mapping[field] = idx
                        break

        # attached workbook, column E holds that value; in the second table its
        # header is mistyped as a year, so the surrounding columns identify it.
        total_distance_columns = [
            idx
            for idx, cell in enumerate(normalized_cells)
            if "totaldistance" in cell or "sectiondistance" in cell
        ]
        if total_distance_columns:
            mapping["distance"] = total_distance_columns[0]
        else:
            has_distance_covered = any("distancecovered" in cell for cell in normalized_cells)
            has_difference = any(cell in {"difference", "diffrence"} for cell in normalized_cells)
            if has_distance_covered and has_difference and len(row_values) >= 5:
                mapping["distance"] = 4

        return mapping

    def _has_route_header(self, mapping: Dict[str, int]) -> bool:
        return all(field in mapping for field in ["start_point", "end_point", "road_type"])

    def _row_from_mapping(
        self,
        values,
        mapping: Dict[str, int],
        serial_counter: int,
        sheet_name: str,
        source_excel_row: int,
        section_number: int,
    ) -> Optional[Dict]:
        start = self._cell(values, mapping.get("start_point"))
        end = self._cell(values, mapping.get("end_point"))
        road_type = self._cell(values, mapping.get("road_type"))

        if not start or not end or not road_type:
            return None
        if self._normalize_header(start) == "total":
            return None

        serial = self._cell(values, mapping.get("serial"))
        if serial in (None, ""):
            serial = serial_counter

        distance = self._cell(values, mapping.get("distance"))

        return {
            "serial": serial,
            "start_point": start,
            "end_point": end,
            "road_type": road_type,
            "distance": distance,
            "source_sheet": sheet_name,
            "source_excel_row": source_excel_row,
            "source_section": section_number,
            "direction": "A to B" if section_number == 1 else "B to A",
        }

    def _auto_detect_columns(self) -> Dict[str, str]:
        """Auto-detect column names from the current dataframe."""
        if self.df is None:
            return {}

        mapping: Dict[str, str] = {}
        excel_columns = [str(col).strip() for col in self.df.columns]

        for field, variations in self.COLUMN_VARIATIONS.items():
            for column in excel_columns:
                normalized_column = self._normalize_header(column)
                if normalized_column == field.replace("_", ""):
                    mapping[field] = column
                    break
                if any(normalized_column == self._normalize_header(v) for v in variations):
                    mapping[field] = column
                    break

        return mapping

    def set_column_mapping(self, mapping: Dict[str, str]) -> None:
        self.column_mapping = mapping

    def get_column_options(self) -> List[str]:
        if self.df is None:
            return []
        return [str(col) for col in self.df.columns]

    def get_data(self) -> pd.DataFrame:
        """Return a dataframe with standardized route columns."""
        if self.df is None:
            return pd.DataFrame()

        result_df = pd.DataFrame()
        for field, column_name in self.column_mapping.items():
            if column_name in self.df.columns:
                result_df[field] = self.df[column_name]

        for extra in ["source_sheet", "source_excel_row", "source_section", "direction"]:
            if extra in self.df.columns and extra not in result_df.columns:
                result_df[extra] = self.df[extra]

        return result_df

    def get_row_by_index(self, index: int) -> Optional[Dict]:
        data = self.get_data()
        if data.empty or index >= len(data):
            return None
        return data.iloc[index].to_dict()

    def is_configured(self) -> bool:
        return all(field in self.column_mapping for field in self.REQUIRED_FIELDS)

    @staticmethod
    def _normalize_header(value) -> str:
        if value is None:
            return ""
        text = str(value).strip().lower()
        return re.sub(r"[^a-z0-9]+", "", text)

    @staticmethod
    def _cell(values, index):
        if index is None or index >= len(values):
            return None
        value = values[index]
        if value is None:
            return None
        if isinstance(value, str):
            value = value.strip()
        return value

    @staticmethod
    def _is_blank_row(values) -> bool:
        return all(value is None or str(value).strip() == "" for value in values)
