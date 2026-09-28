"""
Logger Module
Handles safe logging of rename operations and undo functionality.
"""

import json
import os
from datetime import datetime
from typing import List, Dict, Optional, Tuple
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment


class RenameLogger:
    """Logs rename operations and provides undo functionality."""

    def __init__(self, folder_path: str):
        """
        Initialize the logger.

        Args:
            folder_path: Path where log files will be stored
        """
        self.folder_path = folder_path
        self.json_log_path = os.path.join(folder_path, 'rename_log.json')
        self.xlsx_log_path = os.path.join(folder_path, 'rename_log.xlsx')
        self.log_data = []

    def add_entry(self, old_name: str, new_name: str, status: str, excel_row: int = -1,
                  matched_data: Dict = None, reason: str = "") -> Dict:
        """
        Add an entry to the log.

        Args:
            old_name: Original filename
            new_name: New filename
            status: Status (success, skipped, conflict, not_found, error)
            excel_row: Matching Excel row index
            matched_data: Dictionary of matched Excel data
            reason: Additional reason/notes
        """
        entry = {
            'timestamp': datetime.now().isoformat(),
            'old_name': old_name,
            'new_name': new_name,
            'status': status,
            'excel_row': excel_row,
            'matched_data': matched_data or {},
            'reason': reason
        }
        self.log_data.append(entry)
        return entry

    def save_logs(self) -> Tuple[bool, str]:
        """
        Save logs to JSON and Excel files.

        Returns:
            Tuple of (success: bool, message: str)
        """
        try:
            # Save JSON log
            self._save_json_log()

            # Save Excel log
            self._save_excel_log()

            return True, f"Logs saved: {self.json_log_path} and {self.xlsx_log_path}"

        except Exception as e:
            return False, f"Error saving logs: {str(e)}"

    def _save_json_log(self) -> None:
        """Save log data to JSON file."""
        with open(self.json_log_path, 'w', encoding='utf-8') as f:
            json.dump(self.log_data, f, indent=2, ensure_ascii=False, default=self._json_default)

    def _save_excel_log(self) -> None:
        """Save log data to Excel file."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Rename Log"

        headers = ['Timestamp', 'Old Name', 'New Name', 'Status', 'Serial Number', 'Excel Row',
                   'Sheet', 'Source Row', 'Direction', 'Start Point', 'End Point',
                   'Road Type', 'Distance', 'Reason']
        ws.append(headers)

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Color codes for status
        status_colors = {
            'success': "92D050",  # Green
            'conflict': "FFC7CE",  # Red
            'not_found': "FFEB9C",  # Yellow
            'skipped': "D9E1F2",   # Light blue
            'planned': "D9EAD3",   # Pale green
            'unchanged': "D9E1F2", # Light blue
            'error': "F4B084"      # Orange
        }

        # Data rows
        for entry in self.log_data:
            matched = entry.get('matched_data', {})
            row = [
                entry['timestamp'],
                entry['old_name'],
                entry['new_name'],
                entry['status'],
                matched.get('serial', ''),
                entry['excel_row'] if entry['excel_row'] >= 0 else '',
                matched.get('source_sheet', ''),
                matched.get('source_excel_row', ''),
                matched.get('direction', ''),
                matched.get('start_point', ''),
                matched.get('end_point', ''),
                matched.get('road_type', ''),
                matched.get('distance', ''),
                entry['reason']
            ]
            ws.append(row)

            # Color status cell
            status_cell = ws.cell(row=ws.max_row, column=4)
            status = entry['status']
            if status in status_colors:
                status_cell.fill = PatternFill(start_color=status_colors[status], 
                                               end_color=status_colors[status], 
                                               fill_type="solid")

        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 60
        ws.column_dimensions['C'].width = 60
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 22
        ws.column_dimensions['H'].width = 12
        ws.column_dimensions['I'].width = 12
        ws.column_dimensions['J'].width = 30
        ws.column_dimensions['K'].width = 30
        ws.column_dimensions['L'].width = 18
        ws.column_dimensions['M'].width = 12
        ws.column_dimensions['N'].width = 45

        wb.save(self.xlsx_log_path)

    def load_from_json(self, file_path: str) -> Tuple[bool, str]:
        """
        Load log data from JSON file.

        Args:
            file_path: Path to JSON log file

        Returns:
            Tuple of (success: bool, message: str)
        """
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                self.log_data = json.load(f)
            return True, f"Loaded {len(self.log_data)} entries from log"
        except Exception as e:
            return False, f"Error loading log: {str(e)}"

    def get_undo_list(self) -> List[Tuple[str, str]]:
        """
        Get list of files to undo (reverse of successful renames).

        Returns:
            List of (current_name, original_name) tuples
        """
        undo_list = []
        for entry in reversed(self.log_data):
            if entry['status'] == 'success':
                undo_list.append((entry['new_name'], entry['old_name']))
        return undo_list

    def clear_logs(self) -> None:
        """Clear all log data."""
        self.log_data = []

    @staticmethod
    def _json_default(value):
        """Convert pandas/numpy scalar values to plain JSON-safe values."""
        if hasattr(value, "item"):
            return value.item()
        return str(value)

    def get_summary(self) -> Dict:
        """
        Get summary statistics of the log.

        Returns:
            Dictionary with summary info
        """
        summary = {
            'total_entries': len(self.log_data),
            'successful': 0,
            'conflicts': 0,
            'not_found': 0,
            'skipped': 0,
            'planned': 0,
            'errors': 0
        }

        for entry in self.log_data:
            status = entry['status']
            if status == 'success':
                summary['successful'] += 1
            elif status == 'conflict':
                summary['conflicts'] += 1
            elif status == 'not_found':
                summary['not_found'] += 1
            elif status == 'skipped':
                summary['skipped'] += 1
            elif status == 'planned':
                summary['planned'] += 1
            elif status == 'error':
                summary['errors'] += 1

        return summary
